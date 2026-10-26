"""Bounded, local data profiling. No model or uploaded code is executed."""
import csv
import io
import json
import math
import sqlite3
import statistics
import time
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS = 2000
MAX_COLUMNS = 100
MAX_TABLES = 20
MAX_CELLS_TEXT = 120_000


def quote_identifier(value):
    return '"' + value.replace('"', '""') + '"'


def missing(value):
    return value is None or isinstance(value, str) and not value.strip()


def value_type(value):
    if missing(value):
        return "missing"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number" if math.isfinite(value) else "invalid_number"
    if isinstance(value, (date, datetime)):
        return "date"
    text = str(value).strip()
    try:
        return "number" if math.isfinite(float(text)) else "invalid_number"
    except (ValueError, OverflowError):
        pass
    try:
        date.fromisoformat(text)
        return "date"
    except ValueError:
        return "text"


def validate_rules(rules):
    if not isinstance(rules, dict) or set(rules) - {"required", "types", "ranges"}:
        raise ValueError("Règles : objet avec required, types et ranges uniquement.")
    required, types, ranges = rules.get("required", []), rules.get("types", {}), rules.get("ranges", {})
    if not isinstance(required, list) or len(required) > 100 or not all(isinstance(x, str) and len(x) <= 200 for x in required):
        raise ValueError("required doit être une liste de colonnes.")
    if not isinstance(types, dict) or len(types) > 100 or not all(isinstance(k, str) and v in ("number", "date", "text") for k, v in types.items()):
        raise ValueError("types attend des types number, date ou text.")
    if not isinstance(ranges, dict) or len(ranges) > 100:
        raise ValueError("ranges doit être un objet par colonne.")
    for name, bounds in ranges.items():
        if not isinstance(name, str) or not isinstance(bounds, dict) or not bounds or set(bounds) - {"min", "max"}:
            raise ValueError("Chaque plage attend min et/ou max.")
        if any(type(x) not in (int, float) or not math.isfinite(x) for x in bounds.values()):
            raise ValueError("Bornes numériques finies requises.")
        if bounds.get("min", -math.inf) > bounds.get("max", math.inf):
            raise ValueError("min doit être inférieur ou égal à max.")
    return rules


def profile_table(name, columns, rows, truncated=False, rules=None, structural=None):
    rules = validate_rules(rules or {})
    if not rows:
        normalized = []
    else:
        normalized = [list(row[:len(columns)]) + [None] * max(0, len(columns) - len(row)) for row in rows]
    total_cells = len(rows) * len(columns)
    missing_cells = sum(missing(v) for row in normalized for v in row)
    fingerprints = [json.dumps(row, default=str, ensure_ascii=False) for row in normalized]
    duplicates = len(rows) - len(set(fingerprints))
    all_columns = []
    inconsistent = observed = 0
    checks = failures = 0
    violations = []
    for col in set(rules.get("required", [])) | set(rules.get("types", {})) | set(rules.get("ranges", {})):
        if col not in columns:
            checks += 1
            failures += 1
            violations.append({"column": col, "rule": "colonne absente", "count": 1})
    for i, column in enumerate(columns):
        values = [row[i] for row in normalized]
        types = Counter(value_type(v) for v in values if not missing(v))
        nonempty = sum(types.values())
        majority = types.most_common(1)[0][1] if types else 0
        inconsistent += nonempty - majority
        observed += nonempty
        numeric = [(n + 2, float(v)) for n, v in enumerate(values) if value_type(v) == "number"]
        outliers = []
        if len(numeric) >= 4:
            q1, _, q3 = statistics.quantiles([v for _, v in numeric], n=4, method="inclusive")
            iqr = q3 - q1
            outliers = [line for line, v in numeric if v < q1 - 1.5 * iqr or v > q3 + 1.5 * iqr]
        all_columns.append({"name": column, "missing": sum(missing(v) for v in values), "types": dict(types),
                            "inconsistent_types": nonempty - majority, "outliers": len(outliers),
                            "outlier_rows": outliers[:10]})
        if column in rules.get("required", []):
            checks += len(values)
            count = sum(missing(v) for v in values)
            failures += count
            violations.append({"column": column, "rule": "obligatoire", "count": count})
        if column in rules.get("types", {}):
            present = [v for v in values if not missing(v)]
            checks += len(present)
            count = sum(value_type(v) != rules["types"][column] for v in present)
            failures += count
            violations.append({"column": column, "rule": "type " + rules["types"][column], "count": count})
        if column in rules.get("ranges", {}):
            bounds = rules["ranges"][column]
            present = [v for v in values if not missing(v)]
            checks += len(present)
            count = sum(value_type(v) != "number" or not bounds.get("min", -math.inf) <= float(v) <= bounds.get("max", math.inf) for v in present)
            failures += count
            violations.append({"column": column, "rule": "plage numérique", "count": count})
    return {"table": name, "columns": all_columns, "rows_analyzed": len(rows), "truncated": truncated,
            "missing_cells": missing_cells, "duplicate_rows": duplicates,
            "completeness_pct": round(100 * (1 - missing_cells / total_cells), 2) if total_cells else None,
            "uniqueness_pct": round(100 * (1 - duplicates / len(rows)), 2) if rows else None,
            "type_consistency_pct": round(100 * (1 - inconsistent / observed), 2) if observed else None,
            "rule_compliance_pct": round(100 * (1 - failures / checks), 2) if checks else None,
            "rule_checks": checks, "rule_failures": failures, "violations": violations,
            "structural_issues": structural or [], "rules": rules,
            "governance": "Contrôles techniques sur l’échantillon ; conformité réglementaire, droits d’accès et finalité non évalués.",
            "anomaly_method": "Valeurs numériques hors [Q1 − 1,5×IQR ; Q3 + 1,5×IQR], à vérifier ; ce ne sont pas des erreurs avérées."}


def tabular_files(name, data):
    ext = Path(name).suffix.lower()
    if ext == ".csv":
        text = data.decode("utf-8-sig")
        if len(text) > MAX_CELLS_TEXT or "\x00" in text:
            raise ValueError("CSV trop long ou invalide (UTF-8, 120 000 caractères).")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect=dialect))
        tables = [("donnees", rows)]
    elif ext == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist()) > 1000 or sum(x.file_size for x in z.infolist()) > 24 * 1024 * 1024:
                raise ValueError("Archive XLSX trop volumineuse.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            tables = []
            count = 0
            for sheet in book:
                if len(tables) >= MAX_TABLES or (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError("XLSX limité à 20 feuilles, 2000 lignes de données et 100 colonnes.")
                rows = list(sheet.iter_rows(max_row=min(sheet.max_row or MAX_ROWS + 2, MAX_ROWS + 2), max_col=min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS + 1), values_only=True))
                while rows and all(v is None for v in rows[-1]):
                    rows.pop()
                count += max(0, len(rows) - 1)
                if count > MAX_ROWS:
                    raise ValueError("2000 lignes de données maximum dans le classeur.")
                if rows:
                    tables.append((sheet.title, rows))
        finally:
            book.close()
    else:
        raise ValueError("Import de données : SQLite, CSV ou XLSX requis.")
    if not tables:
        raise ValueError("Aucune table exploitable.")
    result = []
    total = 0
    for name, rows in tables:
        if not rows or not rows[0] or len(rows[0]) > MAX_COLUMNS or len(rows) > MAX_ROWS + 1:
            raise ValueError("En-tête requis ; 2000 lignes et 100 colonnes maximum.")
        columns, issues = [], []
        for i, header in enumerate(rows[0]):
            original = str(header).strip() if header is not None else ""
            candidate = original[:180] or f"colonne_{i + 1}"
            while candidate in columns:
                candidate += "_bis"
            if candidate != original:
                issues.append(f"En-tête vide, tronqué ou répété à la colonne {i + 1} ; nom SQL : {candidate}")
            columns.append(candidate)
        values = []
        for number, row in enumerate(rows[1:], 2):
            if len(row) != len(columns):
                issues.append(f"Ligne {number} : {len(row)} cellules au lieu de {len(columns)}")
            values.append([str(v) if isinstance(v, (date, datetime)) else v for v in row[:len(columns)]] + [None] * max(0, len(columns) - len(row)))
        total += len(json.dumps(values, default=str, ensure_ascii=False))
        if total > MAX_CELLS_TEXT:
            raise ValueError("Données extraites supérieures à 120 000 caractères.")
        result.append((name, columns, values, issues[:100]))
    return result


def build_dataset(path, name, data):
    """Store a private immutable analysis copy; never write into uploaded SQLite."""
    ext = Path(name).suffix.lower()
    if ext in {".db", ".sqlite", ".sqlite3"}:
        if not data.startswith(b"SQLite format 3\x00"):
            raise ValueError("Signature SQLite invalide.")
        path.write_bytes(data)
    else:
        tables = tabular_files(name, data)
        db = sqlite3.connect(path)
        try:
            with db:
                for i, (_, columns, rows, _) in enumerate(tables):
                    table = f"table_{i + 1}"
                    db.execute(f"CREATE TABLE {quote_identifier(table)} ({','.join(quote_identifier(c) for c in columns)})")
                    db.executemany(f"INSERT INTO {quote_identifier(table)} VALUES ({','.join('?' for _ in columns)})", rows)
        finally:
            db.close()
    reports, chunks = inspect_dataset(path)
    if ext not in {".db", ".sqlite", ".sqlite3"}:
        for report, (name, _, _, issues) in zip(reports, tables):
            report["original_name"] = name
            report["structural_issues"] = issues
    return reports, chunks


def inspect_dataset(path, rules=None):
    rules = validate_rules(rules or {})
    db = sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro&immutable=1", uri=True)
    deadline = time.monotonic() + 3
    db.set_progress_handler(lambda: int(time.monotonic() > deadline), 1000)
    db.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, 200_000)
    db.execute("PRAGMA trusted_schema=OFF")
    # Refuse all functions, writes, ATTACH and pragmas while reading user schemas.
    allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ}
    db.set_authorizer(lambda action, *args: sqlite3.SQLITE_OK if action in allowed else sqlite3.SQLITE_DENY)
    try:
        names = [(name, ddl) for name, ddl in db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name LIMIT 101") if not name.startswith("sqlite_")]
        if not names or len(names) > MAX_TABLES:
            raise ValueError("Base limitée à 1–20 tables ordinaires.")
        reports, chunks = [], []
        size = 0
        source_size = 0
        for name, ddl in names:
            if not ddl or re_virtual(ddl):
                raise ValueError("Les tables virtuelles ne sont pas autorisées.")
            cur = db.execute(f"SELECT * FROM {quote_identifier(name)} LIMIT {MAX_ROWS + 1}")
            columns = [c[0] for c in cur.description]
            if len(columns) > MAX_COLUMNS:
                raise ValueError("100 colonnes maximum par table.")
            rows = cur.fetchall()
            truncated = len(rows) > MAX_ROWS
            rows = rows[:MAX_ROWS]
            if any(isinstance(v, bytes) for row in rows for v in row):
                raise ValueError("Les colonnes BLOB ne sont pas prises en charge.")
            size += len(json.dumps(rows, default=str, ensure_ascii=False))
            if size > MAX_CELLS_TEXT:
                raise ValueError("Échantillon supérieur à 120 000 caractères ; importer une base expurgée plus petite.")
            reports.append(profile_table(name, columns, rows, truncated, rules))
            chunks.append((f"table {name}, colonnes", json.dumps(columns, ensure_ascii=False)))
            for i, row in enumerate(rows):
                text = json.dumps(dict(zip(columns, row)), default=str, ensure_ascii=False)
                if source_size + len(text) > MAX_CELLS_TEXT:
                    break
                source_size += len(text)
                for j in range(0, len(text), 1600):
                    chunks.append((f"table {name}, enregistrement échantillonné {i + 1}, extrait {j // 1600 + 1}", text[j:j + 1600]))
        return reports, chunks
    finally:
        db.close()


def re_virtual(ddl):
    import re
    return bool(re.match(r"\s*CREATE\s+VIRTUAL\s+TABLE", ddl, re.I))
